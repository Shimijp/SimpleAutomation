import tkinter
import pandas as pd
from tkinter import filedialog
from docx import Document


excel_filetypes = [
    ("Excel files", "*.xlsx *.xls *.xlsm *.xlsb"),  # Adjust extensions as needed
    ("All files", "*.*")
]
tkinter.Tk().withdraw() # prevents an empty tkinter window from appearing

folder_path = filedialog.askopenfilename(initialdir = "/",
                                          title = "Select a File",
                                          filetypes = excel_filetypes)

df = pd.read_excel(folder_path)
print(df.head)
df1 = df.loc[:, 'Full Name':'Case Num'].copy()
print(df1.head())
def fill_template(template_path, output_path, replacements):
    # Load the template
    doc = Document(template_path)

    # Loop through each paragraph and replace placeholders
    for paragraph in doc.paragraphs:
        for key, value in replacements.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)


    doc.save(output_path)

def create_dict(df):
    list = []
    for row in df.itertuples(index=True,name=None):
        replacements = {
            '[Full Name]': str(row[1]),
            '[.D.I]': str(row[2]),
            '[Date Sent]': str(row[3]),
            '[Case Num]': str(row[4])
        }
        list.append(replacements)
    return list

cases = create_dict(df1)
my_template_path = 'template.docx'
output_dir = filedialog.askdirectory(title="Select Output Directory")
for idx, case in enumerate(cases):
    output_file_path = f"{output_dir}/case_{idx + 1}.docx"
    fill_template(my_template_path, output_file_path, case)
    print(f"Document created: {output_file_path}")

